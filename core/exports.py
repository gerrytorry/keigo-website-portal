"""Isolated canonical profile -> template mapping. Never reads original personal workbook at runtime."""
from pathlib import Path
from io import BytesIO
from datetime import date
from copy import copy
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.drawing.image import Image as XLImage
from django.conf import settings
from .models import Document
SCALARS={'full_name':'E5','name_kana':'E4','gender':'M5','birth_date':'E6','birth_place':'M7','address':'E8','religion':'E10','phone':'E11','marital_status':'S11','health':'S9','blood_type':'U9','height':'O29','weight':'T29','smoking':'O31','alcohol':'T31','tattoo':'O33','passport':'T33','study_months':'O27','japanese_certificate':'T27','strengths':'G36','weaknesses':'G37','hobbies':'G38'}
CHOICES={'M':'男 Pria','F':'女 Wanita','SINGLE':'未婚 Belum menikah','MARRIED':'既婚 Menikah','DIVORCED':'Cerai'}
def safe(value):
 if value is None:return ''
 if isinstance(value,str) and value.startswith(('=','+','-','@')):return "'"+value
 return value
def month_parts(value):
 if not value:return ('','')
 d=date.fromisoformat(str(value)[:10]);return d.year,d.month

def make_xlsx(version):
 data=version.data;w=openpyxl.load_workbook(settings.BASE_DIR/'templates/cv/keigo.xlsx');s=w['Form CV']
 s['M2']=version.created_at.date();s['M2'].number_format='yyyy-mm-dd'
 for field,cell in SCALARS.items():
  value=data.get(field)
  if field=='birth_date' and value:value=date.fromisoformat(value).strftime('%Y年 %m月 %d日')
  elif isinstance(value,bool):value='有 Ya' if value else '無 Tidak'
  elif field in ['gender','marital_status']:value=CHOICES.get(value,value or '')
  elif field=='height' and value is not None:value=str(value)+' cm'
  elif field=='weight' and value is not None:value=str(value)+' kg'
  elif field=='study_months' and value is not None:value=str(value)+' bulan'
  s[cell]=safe(value)
 if data.get('birth_date'):
  born=date.fromisoformat(data['birth_date']);today=version.created_at.date();s['L6']=today.year-born.year-((today.month,today.day)<(born.month,born.day))
 overflow=[]
 for i,row in enumerate(data.get('educations',[])):
  values=[*month_parts(row['start_date']),*month_parts(row.get('end_date')),{'SD':'小学校','SMP':'中学校','SMA':'高校','SMK':'職業高校'}.get(row['level'],row['level']),row['institution'],row.get('major','')]
  if i<3:
   for col,val in zip(['B','C','D','F','G','H','Q'],values):s[f'{col}{14+i}']=safe(val)
  else:overflow.append(('Pendidikan',row))
 for i,row in enumerate(data.get('work_experiences',[])):
  values=[*month_parts(row['start_date']),*month_parts(row.get('end_date')),row['company'],row.get('industry',''),row.get('description',''),row.get('leaving_reason','')]
  if i<5:
   for col,val in zip(['B','C','D','F','G','L','P','T'],values):s[f'{col}{19+i}']=safe(val)
  else:overflow.append(('Pengalaman kerja',row))
 for i,row in enumerate(data.get('certificates',[])):
  if i<2:
   yy,mm=month_parts(row['date']);s[f'B{25+i}']=yy;s[f'C{25+i}']=mm;s[f'D{25+i}']=safe(row['name'])
  else:overflow.append(('Sertifikat',row))
 for i,row in enumerate(data.get('family_members',[])):
  if i<7:
   for col,field in [('B','name'),('G','relationship'),('H','age'),('J','occupation')]:s[f'{col}{28+i}']=safe(row.get(field,''))
  else:overflow.append(('Keluarga',row))
 for name in ['Lanjutan','Tambahan','Dokumen']:
  sh=w.create_sheet(name);sh.sheet_view.showGridLines=False;sh.append(['LPK KEIGO INDONESIA BAGUS','CV versi '+str(version.number)]);sh.append(['Bagian','Field','Nilai']);sh.freeze_panes='A4';sh.column_dimensions['A'].width=25;sh.column_dimensions['B'].width=30;sh.column_dimensions['C'].width=75
 for title,row in overflow:
  for key,value in row.items():
   if key not in ['id','student_id','created_at','updated_at']:w['Lanjutan'].append([title,key,safe(value)])
 for field in ['email','home_phone','eyesight','dominant_hand','english_level','other_languages','favorite_subject','aspirations']:
  w['Tambahan'].append(['Informasi tambahan',field,safe(data.get(field,''))])
 for row in data.get('japan_history',[]):
  for k,val in row.items():
   if k not in ['id','student_id','created_at','updated_at']:w['Tambahan'].append(['Riwayat Jepang',k,safe(val)])
 retained=[];placed=set()
 for entry in sorted(data.get('documents',[]),key=lambda d:d['id'],reverse=True):
  d=Document.objects.filter(pk=entry['id'],student=version.student,digest=entry['digest']).first()
  if not d:raise ValueError('Dokumen snapshot tidak ditemukan')
  w['Dokumen'].append([d.get_category_display(),safe(d.label),safe(d.original_name)])
  if d.category not in placed and d.category in ['PHOTO','LANGUAGE','SKILL'] and Path(d.file.name).suffix.lower() in ['.jpg','.jpeg','.png']:
   try:
    
    with d.file.open('rb') as file:blob=BytesIO(file.read())
    placed.add(d.category);retained.append(blob);im=XLImage(blob);box=(140,170) if d.category=='PHOTO' else (270,210);scale=min(box[0]/im.width,box[1]/im.height);im.width*=scale;im.height*=scale;s.add_image(im,{'PHOTO':'Q3','LANGUAGE':'B41','SKILL':'K41'}[d.category])
   except OSError:raise ValueError('Dokumen gambar tidak tersedia untuk ekspor')
  elif d.category not in placed and d.category in ['LANGUAGE','SKILL']:
   placed.add(d.category)
   s[{'LANGUAGE':'B42','SKILL':'K42'}[d.category]]=safe('Lampiran PDF: '+d.original_name)
 # Full canonical data ensures long/overflow text is not silently lost by fixed template geometry.
 full=w.create_sheet('Data Lengkap');full.append(['LPK KEIGO INDONESIA BAGUS','CV versi '+str(version.number)]);full.append(['Field','Nilai']);full.column_dimensions['A'].width=32;full.column_dimensions['B'].width=100
 for k,val in data.items():
  if not isinstance(val,list):full.append([k,safe(val)])
  elif k!='documents':
   for index,record in enumerate(val,1):
    for field,value in record.items():
     if field not in ['id','student_id','created_at','updated_at']:full.append([f'{k} {index}: {field}',safe(value)])
 for sh in w:
  if sh.title!='Form CV':
   for row in sh:
    for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='Calibri',size=11)
   for c in sh[1]:c.fill=PatternFill('solid',fgColor='132F50');c.font=Font(name='Calibri',size=12,color='FFFFFF',bold=True)
   for row in range(3,sh.max_row+1):sh.row_dimensions[row].height=42
   sh.page_setup.orientation='landscape';sh.page_setup.paperSize=sh.PAPERSIZE_A4;sh.page_setup.fitToWidth=1;sh.page_setup.fitToHeight=0
  else:
   for coord in SCALARS.values():
    style=copy(sh[coord].alignment);style.wrap_text=True;sh[coord].alignment=style
 out=BytesIO();w.save(out);return out.getvalue()
